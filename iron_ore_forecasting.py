"""
===============================================================================
 IRON ORE PRICE FORECASTING: COST CURVE + VECM MODEL  (NPSC3000 / NPSC3001)
===============================================================================

 What this script does, start to finish
 --------------------------------------
   1. Reads the raw data workbook (NPSC3001_RawData.xlsx). It is opened
      READ-ONLY, and a fingerprint (SHA-256 hash) of the file is checked
      before and after the run to prove it was not changed.
   2. Builds the annual seaborne cost curve (2010-2035) and finds the
      marginal producer and marginal cost (MC) in every year.
   3. Builds the seaborne supply-demand balance table (Appendix B).
   4. Runs ADF and KPSS stationarity tests on the six candidate variables.
   5. Runs Johansen cointegration tests on all 15 VECM specifications: the
      restricted-constant case used in the report, the conventional case for
      comparison, every rank step, and the Reinsel-Ahn small-sample correction.
   6. Estimates every VECM and runs the full diagnostic set: long-run
      coefficients, adjustment speeds, short-run terms, stability, residual
      autocorrelation and normality, and Granger causality.
   7. Backtests every model and benchmark recursively, one and two years ahead,
      with identification checks at every forecast origin, accuracy on the
      full and identified samples, and leave-one-year-out robustness.
   8. Produces the final 2026 and 2027 forecasts (each two years ahead), the
      longer 2028-2035 paths, the demand sensitivity and the comparison with
      the Australian Government forecasts (read from the 'Government' column
      of the Macro Data sheet).
   9. Writes two output files next to this script:
         Iron_Ore_Results.docx    - every result table and chart, in one document
         Iron_Ore_Dashboard.html  - an interactive dashboard (works offline):
                                    year, demand and supply sliders, filters,
                                    click-for-detail on every mine, backtest,
                                    forecasts and all diagnostics

 How to run it
 -------------
   1. Install Python 3.10 or newer, then install the packages:
          pip install pandas numpy scipy statsmodels==0.15.0 openpyxl matplotlib python-docx plotly
      statsmodels 0.15.0 is pinned because Model 2's 2021 backtest forecast is
      numerically sensitive at the earliest origin (see Appendix D of the
      report). Every other number is the same across versions.
      plotly is only used to embed the charting library inside the HTML file
      so the dashboard opens without internet.
   2. Put NPSC3001_RawData.xlsx in the same folder as this script.
   3. Run:   python iron_ore_forecasting.py
      or     python iron_ore_forecasting.py "path/to/NPSC3001_RawData.xlsx"
===============================================================================
"""

# =============================================================================
# 0. IMPORTS
# =============================================================================
import hashlib
import io
import json
import math
import sys
import warnings
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import linalg
from statsmodels.stats.diagnostic import acorr_ljungbox
from statsmodels.stats.stattools import jarque_bera
from statsmodels.tsa.api import VAR
from statsmodels.tsa.arima.model import ARIMA
from statsmodels.tsa.stattools import adfuller, kpss
from statsmodels.tsa.vector_ar.vecm import VECM, coint_johansen, select_order

import matplotlib
matplotlib.use("Agg")                     # draw charts without opening windows
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from matplotlib.patches import Patch
from matplotlib.lines import Line2D

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

warnings.filterwarnings("ignore")         # statsmodels prints many harmless warnings


# =============================================================================
# 1. SETTINGS - every modelling choice lives here
# =============================================================================
HERE = Path(__file__).resolve().parent
RAW_DATA = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "NPSC3001_RawData.xlsx"

OUT_DOCX = HERE / "Iron_Ore_Results.docx"
OUT_HTML = HERE / "Iron_Ore_Dashboard.html"

SAMPLE_START, SAMPLE_END = 2010, 2025     # 16 annual observations used for estimation
LAST_YEAR = 2035                          # the workbook projects supply and demand to 2035
FIRST_ORIGIN = 2019                       # first backtest estimation window = 2010-2019
HORIZON = 2                               # the research question: two years ahead
HORIZONS = (1, 2)                         # the backtest also reports one year ahead
FINAL_TARGETS = (2026, 2027)              # final forecasts (origins 2024 and 2025)
K_AR_DIFF = 1                             # one lagged annual difference (= VAR with 2 lags)
SIG = 0.05                                # 5% significance level throughout
ARIMA_ORDERS = [(0, 1, 0), (1, 1, 0), (0, 1, 1), (1, 1, 1)]   # benchmark ARIMA models tried

# Context numbers shown for comparison only (they do not enter any model)
YTD_2026_PRICE = 101.93                   # 2026 year-to-date average (to October 2026), US$/t

# Producer tiers. An asset belongs to a tier if its name contains one of these keys.
TIER1_KEYS = ["BHP WAIO", "Rio Tinto Pilbara", "Fortescue Hematite", "Fortescue Iron Bridge",
              "Vale Northern", "Vale Southeastern", "Vale Southern", "Vale & BHP Germano",
              "Simandou"]
TIER2_KEYS = ["Hancock Roy Hill", "MinRes Onslow", "MinRes Pilbara Hub", "Hancock Atlas",
              "CITIC Sino", "Kumba", "Champion", "Rio Tinto IOC", "CSN", "Anglo Minas-Rio",
              "LKAB", "SNIM", "Shougang", "Ferrexpo", "Metinvest", "CMP ", "Mount Gibson",
              "Grange"]

# The six candidate variables from Table 1 of the report (all in natural logs)
VARIABLES = {"ln_price": "Iron ore price (P)", "ln_mc": "Marginal cost (MC)",
             "ln_demand": "Seaborne demand (D)", "ln_dxy": "US Dollar Index (DXY)",
             "ln_bdi": "Baltic Dry Index (BDI)", "ln_gdp": "China real GDP (GDP)"}
SHORT = {"ln_price": "P", "ln_mc": "MC", "ln_demand": "D", "ln_dxy": "DXY", "ln_bdi": "BDI",
         "ln_gdp": "GDP"}

# The two VECMs carried through the report
MODEL_1 = ["ln_price", "ln_mc"]
MODEL_2 = ["ln_price", "ln_mc", "ln_demand"]
PRINCIPAL = ["Model 1", "Model 2", "Calibrated Marginal Cost Model", "Marginal Cost Model"]

# Johansen 5% critical values for the restricted-constant case ("case 2"),
# from MacKinnon, Haug & Michelis (1999). Keyed by (number of variables - rank).
CV_TRACE = {1: 9.164546, 2: 20.26184, 3: 35.19275, 4: 54.07904, 5: 76.97277}
CV_MAXEIG = {1: 9.164546, 2: 15.89210, 3: 22.29962, 4: 28.58808, 5: 34.80587}

# Chart colours (same palette as the report)
COLOURS = {1: "#2A4879", 2: "#406E6C", 3: "#AB7A5D", "price": "#6C727F", "mc": "#83525D",
           "demand": "#8C8C8C"}


# =============================================================================
# 2. READ THE RAW DATA WORKBOOK
# =============================================================================
def sha256(path):
    """Fingerprint of the file. If even one cell changed, this would change."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def year_columns(ws, header_row):
    """Find which spreadsheet column holds which year, e.g. {2010: 3, 2011: 4, ...}."""
    cols = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, (int, float)) and 2000 < v < 2060:
            cols[int(v)] = c
    return cols


def read_grid(ws, header_row, first_row=None):
    """Read a 'one row per asset, one column per year' sheet into
    {asset name: {year: value}}. Rows labelled 'Total...' are skipped."""
    yc = year_columns(ws, header_row)
    out = {}
    for r in range(first_row or header_row + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label is None or not str(label).strip():
            continue
        label = str(label).strip()
        if label.lower().startswith("total"):
            continue
        out[label] = {y: ws.cell(row=r, column=c).value for y, c in yc.items()}
    return out


def load_raw_data(path):
    """Read everything the model needs. data_only=True means we read the
    calculated values, not the formulas. The file is never saved."""
    print(f"Reading {path.name} ...")
    wb = load_workbook(path, data_only=True)

    supply = read_grid(wb["Seaborne Supply"], 2)          # Mt wet, by asset
    cfr = read_grid(wb["Final CFR cost"], 3)              # US$/t delivered to Qingdao, by asset

    fob = read_grid(wb["FOB Cost"], 2)                    # US$/wmt mine-gate cost (for the dashboard detail panel)
    freight = read_grid(wb["Freight to Qingdao"], 3, first_row=7)   # US$/wmt freight to Qingdao

    # Descriptive information about each asset (operator, country, port ...)
    ov = wb["Asset overview"]
    meta = {}
    for r in range(3, ov.max_row + 1):
        name = ov.cell(r, 1).value
        if name:
            meta[str(name).strip()] = {k: ov.cell(r, c).value or "" for k, c in
                                       (("operator", 2), ("owner", 3), ("country", 4), ("region", 5),
                                        ("port", 6), ("status", 7))}
    asm = wb["Assumptions"]
    assumptions = {}
    for r in range(3, asm.max_row + 1):
        name = asm.cell(r, 1).value
        if name:
            assumptions[str(name).strip()] = {"grade": asm.cell(r, 3).value, "moisture": asm.cell(r, 4).value,
                                              "royalty": asm.cell(r, 5).value}

    # Demand build: pick out the rows we need by their label
    ws = wb["Demand build"]
    yc = year_columns(ws, 3)
    rows = {str(ws.cell(r, 1).value).strip(): r for r in range(4, ws.max_row + 1)
            if ws.cell(r, 1).value}

    def demand_row(label):
        r = next(rr for lbl, rr in rows.items() if label.lower() in lbl.lower())
        return {y: ws.cell(row=r, column=c).value for y, c in yc.items()}

    demand = {"total": demand_row("Total Seaborne Demand"),
              "china": demand_row("China Seaborne Imports"),
              "other": demand_row("Other Seaborne Demand"),
              "consumption": demand_row("Gross Ore Requirement"),
              "domestic": demand_row("China Domestic Usable"),
              "overland": demand_row("Overland"),
              "stocks": demand_row("Change in Visible"),
              "adjustment": demand_row("Statistical / product-mix")}

    # Macro data: price (62% Fe), DXY, BDI, China real GDP and, where given,
    # the Australian Government's two-year-ahead forecast for that year
    md = wb["Macro Data"]
    headers = [str(md.cell(1, c).value or "").strip().lower() for c in range(1, md.max_column + 1)]
    gov_col = headers.index("government") + 1 if "government" in headers else None
    macro, government = {}, {}
    for r in range(2, md.max_row + 1):
        y = md.cell(r, 1).value
        if y is not None:
            macro[int(y)] = {"price": md.cell(r, 2).value, "dxy": md.cell(r, 3).value,
                             "bdi": md.cell(r, 4).value, "gdp": md.cell(r, 5).value}
            if gov_col and md.cell(r, gov_col).value is not None:
                government[int(y)] = float(md.cell(r, gov_col).value)
    wb.close()
    return {"supply": supply, "cfr": cfr, "fob": fob, "freight": freight, "meta": meta,
            "assumptions": assumptions, "demand": demand, "macro": macro, "government": government}


def tier_of(asset):
    """Tier 1, 2 or 3 for an asset name."""
    if any(asset.startswith(k) or k in asset for k in TIER1_KEYS):
        return 1
    if any(asset.startswith(k) or k in asset for k in TIER2_KEYS):
        return 2
    return 3


# =============================================================================
# 3. THE COST CURVE, THE MARGINAL COST AND THE MARKET BALANCE
# =============================================================================
def cost_curve(D, year):
    """All assets producing in `year`, ranked from cheapest to most expensive
    delivered (CFR Qingdao) cost, with running cumulative supply."""
    rows = []
    for asset, sup in D["supply"].items():
        volume = sup.get(year) or 0.0
        cost = (D["cfr"].get(asset) or {}).get(year) or 0.0
        if volume > 0 and cost > 0:
            m, a = D["meta"].get(asset, {}), D["assumptions"].get(asset, {})
            rows.append({"asset": asset, "tier": tier_of(asset), "supply_mt": float(volume), "cost": float(cost),
                         "fob": (D["fob"].get(asset) or {}).get(year),
                         "freight": (D["freight"].get(asset) or {}).get(year),
                         "grade": a.get("grade"), "moisture": a.get("moisture"), "royalty": a.get("royalty"),
                         **{k: m.get(k, "") for k in ("operator", "owner", "country", "region", "port", "status")}})
    rows.sort(key=lambda r: r["cost"])                    # cheapest first
    cumulative = 0.0
    for r in rows:
        cumulative += r["supply_mt"]
        r["cumulative_mt"] = cumulative
    return rows


def marginal_cost(D, year, demand=None):
    """The marginal producer is the first asset at which cumulative supply
    reaches demand. Its cost is the marginal cost for that year."""
    demand = D["demand"]["total"][year] if demand is None else demand
    for r in cost_curve(D, year):
        if r["cumulative_mt"] >= demand:
            return r["cost"], r["asset"]
    return None, None                                     # supply never reaches demand


def build_dataset(D):
    """One row per year, 2010-2035, with every series the econometrics needs."""
    rec = []
    for y in range(SAMPLE_START, LAST_YEAR + 1):
        curve = cost_curve(D, y)
        mc, asset = marginal_cost(D, y)
        m = D["macro"].get(y, {})
        rec.append({"year": y, "price": m.get("price"), "dxy": m.get("dxy"), "bdi": m.get("bdi"),
                    "gdp": m.get("gdp"), "marginal_cost": mc, "marginal_asset": asset,
                    "demand": D["demand"]["total"][y], "total_supply": curve[-1]["cumulative_mt"],
                    "assets_producing": len(curve)})
    df = pd.DataFrame(rec).set_index("year")
    log = lambda v: math.log(v) if v is not None and not pd.isna(v) and v > 0 else np.nan
    df["ln_price"], df["ln_mc"] = df.price.map(log), df.marginal_cost.map(log)
    df["ln_demand"], df["ln_dxy"] = df.demand.map(log), df.dxy.map(log)
    df["ln_bdi"], df["ln_gdp"] = df.bdi.map(log), df.gdp.map(log)
    return df


def supply_by(D, year, *keys):
    """Total supply of every asset whose name contains any of `keys`."""
    return sum((s.get(year) or 0.0) for n, s in D["supply"].items() if any(k in n for k in keys))


def market_balance(D, df):
    """Appendix B: seaborne demand build, supply by producer, Tier-3 required
    and market balance. Rio Tinto includes IOC (Canada); the Germano joint
    venture is split 50/50 between BHP and Vale; Tier-2 excludes IOC so it is
    not counted twice. Tier-3 = demand - named producers - Tier-2 (negative
    means Tier 1-2 supply alone exceeds demand). Market balance = change in
    Tier-3 required + change in visible inventories."""
    d = D["demand"]
    out = []
    for y in range(SAMPLE_START, LAST_YEAR + 1):
        rio = supply_by(D, y, "Rio Tinto Pilbara", "Rio Tinto IOC")
        bhp = supply_by(D, y, "BHP WAIO") + 0.5 * supply_by(D, y, "Vale & BHP Germano")
        fmg = supply_by(D, y, "Fortescue Hematite", "Fortescue Iron Bridge")
        vale = (supply_by(D, y, "Vale Northern", "Vale Southeastern", "Vale Southern")
                + 0.5 * supply_by(D, y, "Vale & BHP Germano"))
        sim = supply_by(D, y, "Simandou")
        t2 = sum((s.get(y) or 0.0) for n, s in D["supply"].items()
                 if tier_of(n) == 2 and "Rio Tinto IOC" not in n)
        total = d["total"][y]
        tier1 = rio + bhp + fmg + vale + sim
        out.append({"Year": y, "China seaborne demand": d["china"][y],
                    "Total consumption": d["consumption"][y],
                    "Total imports": (d["china"][y] or 0) + (d["overland"][y] or 0),
                    "Change in visible inventories": d["stocks"][y],
                    "Apparent domestic production": d["domestic"][y],
                    "Statistical / product-mix adjustment": d["adjustment"][y],
                    "Overland imports": d["overland"][y], "Other seaborne imports": d["other"][y],
                    "Total seaborne demand": total, "Tier-1 supply": tier1, "Rio Tinto": rio,
                    "BHP": bhp, "FMG": fmg, "Vale": vale, "Simandou": sim, "Tier-2 supply": t2,
                    "Tier-3 supply (required)": total - tier1 - t2,
                    "Marginal cost (US$/t)": df.at[y, "marginal_cost"],
                    "62% Fe price, realised (US$/t)": df.at[y, "price"] if y <= SAMPLE_END else None})
    bal = pd.DataFrame(out).set_index("Year")
    bal.insert(bal.columns.get_loc("Marginal cost (US$/t)"), "Market balance",
               bal["Tier-3 supply (required)"].diff() + bal["Change in visible inventories"])
    return bal


def integrity_checks(df, hash_before):
    """Simple checks that the data behave as expected."""
    rows = [{"Check": "Raw data file fingerprint (SHA-256) before run", "Result": hash_before}]
    for y in df.index:
        rows.append({"Check": f"{y}: cost curve reaches seaborne demand",
                     "Result": "Yes" if df.at[y, "total_supply"] >= df.at[y, "demand"] else "NO"})
    missing = [y for y in range(SAMPLE_START, SAMPLE_END + 1)
               if df.loc[y, ["price", "dxy", "bdi", "gdp", "marginal_cost"]].isna().any()]
    rows.append({"Check": "Missing values in 2010-2025 estimation data",
                 "Result": "None found" if not missing else str(missing)})
    return pd.DataFrame(rows)


# =============================================================================
# 4. STATIONARITY TESTS (ADF and KPSS)
# =============================================================================
def kpss_p(x, reg):
    """KPSS p-value (interpolated from tables, so it is bounded at 0.01 and 0.10)."""
    try:
        return kpss(x, regression=reg, nlags="auto")[1]
    except Exception:
        return np.nan


def unit_root(series, det):
    """ADF (null: unit root) and KPSS (null: stationary), in levels and first
    differences. `det` is 'c' (constant) or 'ct' (constant + trend) for the
    levels test; first differences always use a constant."""
    x = series.dropna().values
    dx = np.diff(x)
    adf_l = adfuller(x, maxlag=2, autolag="AIC", regression=det)
    adf_d = adfuller(dx, maxlag=2, autolag="AIC", regression="c")
    kp_l, kp_d = kpss_p(x, det), kpss_p(dx, "c")
    # I(1) at 5%: unit root in levels (ADF not rejected) and stationary after
    # differencing (ADF rejects and KPSS does not reject)
    i1 = adf_l[1] >= SIG and adf_d[1] < SIG and kp_d >= SIG
    if adf_l[1] >= SIG and kp_l < SIG and i1:
        result = "I(1) - both tests agree"
    elif i1:
        result = "I(1) - suggestive (KPSS does not confirm the level is non-stationary)"
    elif adf_l[1] < SIG and kp_l < SIG:
        result = "Conflicting - ADF says stationary, KPSS says non-stationary"
    else:
        result = "Not I(1) - first difference is not clearly stationary"
    return {"ADF level stat": adf_l[0], "ADF level p": adf_l[1], "ADF level lags": adf_l[2],
            "ADF diff stat": adf_d[0], "ADF diff p": adf_d[1], "ADF diff lags": adf_d[2],
            "KPSS level p": kp_l, "KPSS diff p": kp_d, "I(1) at 5%": i1, "Result": result}


def stationarity_table(hist):
    rows = []
    for col, name in VARIABLES.items():
        for det, label in (("c", "Constant"), ("ct", "Constant + trend")):
            rows.append({"Variable": name, "col": col, "Terms": label, **unit_root(hist[col], det)})
    t = pd.DataFrame(rows)
    # A variable goes forward if it is I(1) under at least one specification
    t["Proceeds to cointegration testing"] = t.groupby("col")["I(1) at 5%"].transform("any")
    return t


# =============================================================================
# 5. JOHANSEN COINTEGRATION TESTS
# =============================================================================
def johansen_case2(Y, k_ar_diff=K_AR_DIFF):
    """Johansen reduced-rank regression with the constant restricted to the
    cointegrating relationship ('case 2'). This matches how the VECM is
    estimated (deterministic='ci'); statsmodels' own coint_johansen does not
    offer this case, so it is written out here.

    Returns the eigenvalues, trace and maximum-eigenvalue statistics for every
    rank step, their 5% critical values and the effective number of observations."""
    T, p = Y.shape
    dY = np.diff(Y, axis=0)
    idx = range(k_ar_diff + 1, T)
    Z0 = np.array([Y[t] - Y[t - 1] for t in idx])                 # this year's change
    Z1 = np.array([np.r_[Y[t - 1], 1.0] for t in idx])            # last year's level + constant
    if k_ar_diff:                                                 # remove the lagged changes
        Z2 = np.array([np.concatenate([dY[t - 1 - j] for j in range(1, k_ar_diff + 1)]) for t in idx])
        resid = lambda A: A - Z2 @ np.linalg.lstsq(Z2, A, rcond=None)[0]
        R0, R1 = resid(Z0), resid(Z1)
    else:
        R0, R1 = Z0, Z1
    n = len(R0)
    S00, S11, S01 = R0.T @ R0 / n, R1.T @ R1 / n, R0.T @ R1 / n
    eig = linalg.eigh(S01.T @ np.linalg.solve(S00, S01), S11, eigvals_only=True)
    eig = np.clip(np.sort(eig)[::-1][:p], 0, 1 - 1e-12)
    trace = [-n * np.sum(np.log(1 - eig[r:])) for r in range(p)]
    maxeig = [-n * np.log(1 - eig[r]) for r in range(p)]
    return {"eig": eig, "trace": trace, "maxeig": maxeig, "n_eff": n,
            "cv_trace": [CV_TRACE[p - r] for r in range(p)],
            "cv_maxeig": [CV_MAXEIG[p - r] for r in range(p)]}


def reinsel_ahn(stats, n_eff, p):
    """Small-sample correction: multiply each statistic by (T - p*k)/T, where k
    is the number of lags in the levels VAR (one lagged difference = 2 lags)."""
    factor = (n_eff - p * (K_AR_DIFF + 1)) / n_eff
    return [s * factor for s in stats], factor


def sequential_rank(stats, cvs):
    """Johansen 'ladder': test rank 0, then rank <= 1, and so on, stopping at
    the first hypothesis that is NOT rejected. That stopping point is the rank."""
    r = 0
    for s, c in zip(stats, cvs):
        if s > c:
            r += 1
        else:
            break
    return r


def rank_label(trace_rank, max_rank, p):
    """Summarise the two tests: the shared rank if they agree, 'Full' if a test
    says every variable is stationary (inconsistent with I(1) data), otherwise
    'Mixed' (the tests disagree)."""
    if trace_rank == max_rank:
        return "Full" if trace_rank == p else str(trace_rank)
    if trace_rank == p or max_rank == p:
        return "Full"
    return "Mixed"


def n_unit_canonical_correlations(Y):
    """Two or more canonical correlations equal to 1 means a rank-1 VECM is not
    identified (the long-run relationship is not unique). A singular matrix
    means the window is too short for the number of variables."""
    try:
        return int(sum(e > 0.9999 for e in johansen_case2(Y)["eig"]))
    except np.linalg.LinAlgError:
        return Y.shape[1]


def identification_status(n_unit):
    if n_unit >= 2:
        return "Not identified"
    if n_unit == 1:
        return "Degenerate (perfect in-sample fit), identified"
    return "Identified"


def spec_name(cols):
    name = " + ".join(SHORT[c] for c in cols)
    return {"P + MC": "Model 1: P + MC", "P + MC + D": "Model 2: P + MC + D"}.get(name, name)


def johansen_all_steps(hist, cols):
    """Every rank step for one specification, for both the restricted-constant
    case (used in the report) and the conventional unrestricted-constant case
    (statsmodels coint_johansen, det_order=0) for comparison."""
    Y = hist[cols].values
    p = len(cols)
    rows = []
    J = johansen_case2(Y)
    tr_c, _ = reinsel_ahn(J["trace"], J["n_eff"], p)
    mx_c, _ = reinsel_ahn(J["maxeig"], J["n_eff"], p)
    conv = coint_johansen(Y, 0, K_AR_DIFF)
    for r in range(p):
        rows.append({"Specification": spec_name(cols), "Null hypothesis": f"r = {r} (trace: r <= {r})",
                     "Eigenvalue": J["eig"][r],
                     "Trace (restricted constant)": J["trace"][r], "Trace 5% CV": J["cv_trace"][r],
                     "Trace corrected (R-A)": tr_c[r],
                     "Max-eig (restricted constant)": J["maxeig"][r], "Max-eig 5% CV": J["cv_maxeig"][r],
                     "Max-eig corrected (R-A)": mx_c[r],
                     "Trace (conventional)": conv.lr1[r], "Trace CV (conventional)": conv.cvt[r, 1],
                     "Max-eig (conventional)": conv.lr2[r], "Max-eig CV (conventional)": conv.cvm[r, 1]})
    return rows


# =============================================================================
# 6. VECM ESTIMATION AND DIAGNOSTICS
# =============================================================================
def fit_vecm(frame, cols):
    """VECM with one lagged difference, rank 1, constant restricted to the
    long-run relationship. Price is listed first so the relationship is
    normalised on price."""
    return VECM(frame[cols].values, k_ar_diff=K_AR_DIFF, coint_rank=1, deterministic="ci").fit()


def system_stability(res):
    """Largest companion-matrix root that is not one of the unit roots implied
    by cointegration. Below 1 = stable; above 1 = explosive."""
    A = res.var_rep
    p, kk = A.shape[1], A.shape[0]
    comp = np.zeros((p * kk, p * kk))
    comp[:p, :] = np.hstack(list(A))
    if kk > 1:
        comp[p:, :-p] = np.eye(p * (kk - 1))
    mods = np.sort(np.abs(np.linalg.eigvals(comp)))[::-1]
    non_unit = sorted(mods, key=lambda m: abs(m - 1))[p - res.coint_rank:]
    return float(max(non_unit)) if non_unit else 0.0


def vecm_diagnostics(hist, cols):
    """Everything about one full-sample VECM, flattened into one row."""
    res = fit_vecm(hist, cols)
    row = {"Specification": spec_name(cols), "Intercept (long run)": -float(res.det_coef_coint[0, 0]),
           "Intercept s.e.": float(res.stderr_det_coef_coint[0, 0])}
    for i, c in enumerate(cols[1:], start=1):                     # long-run coefficients on price
        row[f"Long-run {SHORT[c]}"] = -float(res.beta[i, 0])
        row[f"Long-run {SHORT[c]} s.e."] = float(res.stderr_beta[i, 0])
        row[f"Long-run {SHORT[c]} p"] = float(res.pvalues_beta[i, 0])
    for i, c in enumerate(cols):                                  # adjustment speed of every variable
        row[f"Alpha {SHORT[c]}"] = float(res.alpha[i, 0])
        row[f"Alpha {SHORT[c]} p"] = float(res.pvalues_alpha[i, 0])
    for j, c in enumerate(cols):                                  # short-run terms in the price equation
        row[f"Price eq: d.{SHORT[c]}(t-1)"] = float(np.asarray(res.gamma)[0, j])
        row[f"Price eq: d.{SHORT[c]}(t-1) p"] = float(res.pvalues_gamma[0, j])
    row["Largest non-unit root"] = system_stability(res)
    row["Stable system"] = row["Largest non-unit root"] < 1
    for i, c in enumerate(cols):                                  # residual checks, equation by equation
        e = res.resid[:, i]
        row[f"Ljung-Box (lag 2) p, {SHORT[c]} eq"] = float(acorr_ljungbox(e, lags=[2], return_df=True)["lb_pvalue"].iloc[0])
        row[f"Jarque-Bera p, {SHORT[c]} eq"] = float(jarque_bera(e)[1])
    for test, label in ((lambda: res.test_normality().pvalue, "System normality p"),
                        (lambda: res.test_whiteness(nlags=3).pvalue, "System whiteness (3 lags) p")):
        try:
            row[label] = float(test())
        except Exception:
            row[label] = None
    for j, c in enumerate(cols[1:], start=1):                     # does each variable help predict price?
        try:
            row[f"Granger: {SHORT[c]} -> P, p"] = float(res.test_granger_causality(caused=0, causing=j).pvalue)
        except Exception:
            row[f"Granger: {SHORT[c]} -> P, p"] = None
    return row


def lag_selection(hist, cols):
    """Information criteria for the number of lagged differences (searched 0-2).
    The report uses one; a criterion choosing 2 sits at the edge of the search."""
    sel = select_order(hist[cols].values, maxlags=2, deterministic="ci")
    return {"Specification": spec_name(cols), "AIC": int(sel.aic), "BIC": int(sel.bic),
            "HQIC": int(sel.hqic), "FPE": int(sel.fpe), "Used in report": K_AR_DIFF}


# =============================================================================
# 7. FORECASTING METHODS (each one only sees the training data it is given)
# =============================================================================
def vecm_forecast(train, cols, steps=HORIZON):
    """Estimate on `train` only, then forecast `steps` years ahead (US$/t)."""
    res = fit_vecm(train, cols)
    return [float(np.exp(row[0])) for row in res.predict(steps=steps)]


def var_diff_forecast(train, cols, steps=HORIZON):
    """VAR in first differences (assumes no cointegration). Forecast the
    changes, then add them to the last observed log price."""
    d = train[cols].diff().dropna().values
    res = VAR(d).fit(1)
    f = res.forecast(d[-1:], steps)
    last = train[cols].values[-1, 0]
    return [float(np.exp(last + f[: h + 1, 0].sum())) for h in range(steps)]


def markup(train, how="geometric"):
    """Calibration factor. Geometric (used in the report): exp(mean(ln P - ln MC)).
    Arithmetic (robustness check): mean(P / MC)."""
    ratio = train.price / train.marginal_cost
    return float(np.exp(np.mean(np.log(ratio)))) if how == "geometric" else float(np.mean(ratio))


def arima_forecast(train, steps=HORIZON):
    """Best of a small ARIMA grid on log price (lowest AIC)."""
    s = train.ln_price.values
    best = (np.inf, None, None)
    for order in ARIMA_ORDERS:
        try:
            r = ARIMA(s, order=order).fit()
            if np.isfinite(r.aic) and r.aic < best[0]:
                best = (r.aic, r.forecast(steps=steps), order)
        except Exception:
            pass
    return [float(np.exp(v)) for v in best[1]], str(best[2])


def method_forecast(method, train, steps=HORIZON):
    """Dispatch one method. Structural models use the marginal cost known at
    the origin, held flat (historical forecasts of future MC do not exist)."""
    mc_now = float(train.marginal_cost.iloc[-1])
    if method["kind"] == "vecm":
        return vecm_forecast(train, method["cols"], steps)
    if method["kind"] == "var_diff":
        return var_diff_forecast(train, method["cols"], steps)
    if method["kind"] == "calibrated":
        return [mc_now * markup(train, method["markup"])] * steps
    if method["kind"] == "raw_mc":
        return [mc_now] * steps
    if method["kind"] == "arima":
        return arima_forecast(train, steps)[0]
    if method["kind"] == "random_walk":
        return [float(train.price.iloc[-1])] * steps
    raise ValueError(method)


# =============================================================================
# 8. BACKTEST ENGINE - recursive (expanding window), one and two years ahead
# =============================================================================
def origins(hist):
    """Every year from 2019 that still has at least one later year to test."""
    return [T for T in hist.index if FIRST_ORIGIN <= T < SAMPLE_END]


def run_backtest(hist, methods):
    """For each origin T: cut the data at T (nothing later is visible),
    re-estimate every method, and forecast T+1 and T+2."""
    recs = []
    for T in origins(hist):
        train = hist.loc[:T]
        for name, m in methods.items():
            path = method_forecast(m, train, max(HORIZONS))
            n_unit = n_unit_canonical_correlations(train[m["cols"]].values) if m["kind"] == "vecm" else 0
            for h in HORIZONS:
                target = T + h
                if target > SAMPLE_END:
                    continue
                actual = float(hist.at[target, "price"])
                recs.append({"Method": name, "Group": m["group"], "Horizon": h, "Origin": T,
                             "Target": target, "Estimation period": f"{SAMPLE_START}-{T}",
                             "Effective obs": len(train) - K_AR_DIFF - 1, "Actual": actual,
                             "Forecast": path[h - 1], "Error": path[h - 1] - actual,
                             "Abs % error": abs(path[h - 1] - actual) / actual * 100,
                             "Unit canonical correlations": n_unit,
                             "Identification": identification_status(n_unit)})
    return pd.DataFrame(recs)


def metrics(sub):
    """RMSE, MAE, MAPE and mean error (forecast - actual) for a set of forecasts."""
    e = sub.Forecast.values - sub.Actual.values
    return {"n": len(sub), "RMSE": float(np.sqrt(np.mean(e ** 2))), "MAE": float(np.mean(np.abs(e))),
            "MAPE": float(np.mean(np.abs(e / sub.Actual.values)) * 100), "Mean error": float(np.mean(e))}


def accuracy_tables(bt):
    """Accuracy for every method and horizon, on (a) all forecasts and (b)
    only the targets where Model 1 and Model 2 are both identified."""
    bad = bt[bt.Method.isin(["Model 1", "Model 2"]) & (bt.Identification == "Not identified")]
    drop = set(zip(bad.Horizon, bad.Target))
    ident = bt[[(h, t) not in drop for h, t in zip(bt.Horizon, bt.Target)]]
    rows = []
    for label, frame in (("All forecasts", bt), ("Identified sample (Models 1-2)", ident)):
        for (m, h), sub in frame.groupby(["Method", "Horizon"], sort=False):
            bad_n = int((sub.Identification == "Not identified").sum())
            rows.append({"Sample": label, "Method": m, "Group": sub.Group.iloc[0], "Horizon": h, **metrics(sub),
                         "Non-identified forecasts": bad_n,
                         "Reliable": bad_n <= 1 and np.isfinite(metrics(sub)["RMSE"])})
    return pd.DataFrame(rows)


def leave_one_out(bt, horizon=HORIZON):
    """Drop one target year at a time and re-rank the four principal models by
    RMSE. Shows whether one year drives the ranking."""
    sub = bt[(bt.Horizon == horizon) & bt.Method.isin(PRINCIPAL)]
    rows = []
    for year in sorted(sub.Target.unique()):
        keep = sub[sub.Target != year]
        rmse = {m: metrics(keep[keep.Method == m])["RMSE"] for m in PRINCIPAL}
        order = sorted(rmse, key=rmse.get)
        rows.append({"Horizon": horizon, "Year left out": year, **{f"RMSE {m}": rmse[m] for m in PRINCIPAL},
                     "Best model": order[0]})
    return pd.DataFrame(rows)


def origin_johansen(hist, cols):
    """Repeat the Johansen test at every forecast origin, to show how the
    evidence for cointegration changes as the sample grows."""
    rows = []
    for T in origins(hist) + [SAMPLE_END]:
        Y = hist.loc[:T, cols].values
        p = len(cols)
        try:
            J = johansen_case2(Y)
            tr_c, _ = reinsel_ahn(J["trace"], J["n_eff"], p)
            rows.append({"Specification": spec_name(cols), "Data to": T, "Effective obs": J["n_eff"],
                         "Trace r=0": J["trace"][0], "5% CV": J["cv_trace"][0],
                         "Trace rank": sequential_rank(J["trace"], J["cv_trace"]),
                         "Max-eig rank": sequential_rank(J["maxeig"], J["cv_maxeig"]),
                         "Corrected trace rank": sequential_rank(tr_c, J["cv_trace"]),
                         "Identification": identification_status(n_unit_canonical_correlations(Y))})
        except np.linalg.LinAlgError:
            rows.append({"Specification": spec_name(cols), "Data to": T,
                         "Identification": "Not identified (singular)"})
    return pd.DataFrame(rows)


# =============================================================================
# 9. RUN EVERYTHING
# =============================================================================
def run():
    hash_before = sha256(RAW_DATA)
    D = load_raw_data(RAW_DATA)
    df = build_dataset(D)
    hist = df.loc[SAMPLE_START:SAMPLE_END].copy()

    # ---- Cost curves, Appendix B and data checks ------------------------------
    balance = market_balance(D, df)
    curves = {y: cost_curve(D, y) for y in range(SAMPLE_START, LAST_YEAR + 1)}
    checks = integrity_checks(df, hash_before)

    # ---- Table A1: stationarity ------------------------------------------------
    print("Stationarity tests ...")
    stat = stationarity_table(hist)
    qualifying = [c for c in VARIABLES if stat.loc[stat.col == c, "Proceeds to cointegration testing"].any()]
    explanatory = [c for c in qualifying if c != "ln_price"]

    # ---- Tables A2 and A3 + full diagnostics, every specification ---------------
    print("Johansen tests and VECM diagnostics for every specification ...")
    specs = [["ln_price", *combo] for n in range(1, len(explanatory) + 1)
             for combo in combinations(explanatory, n)]
    summary, steps, diags = [], [], []
    for cols in specs:
        p = len(cols)
        J = johansen_case2(hist[cols].values)
        tr_c, factor = reinsel_ahn(J["trace"], J["n_eff"], p)
        mx_c, _ = reinsel_ahn(J["maxeig"], J["n_eff"], p)
        r_tr, r_mx = sequential_rank(J["trace"], J["cv_trace"]), sequential_rank(J["maxeig"], J["cv_maxeig"])
        r_trc, r_mxc = sequential_rank(tr_c, J["cv_trace"]), sequential_rank(mx_c, J["cv_maxeig"])
        res = fit_vecm(hist, cols)
        alpha, alpha_p = float(res.alpha[0, 0]), float(res.pvalues_alpha[0, 0])
        summary.append({
            "Specification": spec_name(cols), "Variables": p, "Effective obs": J["n_eff"],
            "R-A factor": factor,
            "Trace r=0": J["trace"][0], "Trace CV r=0": J["cv_trace"][0], "Corrected trace r=0": tr_c[0],
            "Max-eig r=0": J["maxeig"][0], "Max-eig CV r=0": J["cv_maxeig"][0], "Corrected max-eig r=0": mx_c[0],
            "Trace rank": r_tr, "Max-eig rank": r_mx, "Corrected trace rank": r_trc, "Corrected max-eig rank": r_mxc,
            "Result (unadjusted -> corrected)": f"{rank_label(r_tr, r_mx, p)} -> {rank_label(r_trc, r_mxc, p)}",
            "Passes cointegration (unadjusted)": r_tr == 1 and r_mx == 1,
            "Full-sample identification": identification_status(n_unit_canonical_correlations(hist[cols].values)),
            "Price alpha": alpha, "Alpha p-value": alpha_p,
            "Alpha assessment": ("Negative; significant" if alpha < 0 and alpha_p < SIG else
                                 "Positive; significant" if alpha > 0 and alpha_p < SIG else "Not significant"),
            "cols": cols})
        steps += johansen_all_steps(hist, cols)
        diags.append(vecm_diagnostics(hist, cols))
    joh = pd.DataFrame(summary)
    steps, diags = pd.DataFrame(steps), pd.DataFrame(diags)
    lags = pd.DataFrame([lag_selection(hist, MODEL_1), lag_selection(hist, MODEL_2)])

    # ---- Backtest: every VECM + structural models + benchmarks --------------------
    print("Backtesting every model (one and two years ahead) ...")
    methods = {"Model 1": {"kind": "vecm", "cols": MODEL_1, "group": "Principal"},
               "Model 2": {"kind": "vecm", "cols": MODEL_2, "group": "Principal"},
               "Calibrated Marginal Cost Model": {"kind": "calibrated", "markup": "geometric", "cols": [], "group": "Principal"},
               "Marginal Cost Model": {"kind": "raw_mc", "cols": [], "group": "Principal"}}
    for cols in specs:
        if cols not in (MODEL_1, MODEL_2):
            methods[spec_name(cols)] = {"kind": "vecm", "cols": cols, "group": "Other VECM"}
    methods.update({
        "Calibrated (arithmetic markup)": {"kind": "calibrated", "markup": "arithmetic", "cols": [], "group": "Benchmark"},
        "VAR in differences (P + MC)": {"kind": "var_diff", "cols": MODEL_1, "group": "Benchmark"},
        "ARIMA (best AIC)": {"kind": "arima", "cols": [], "group": "Benchmark"},
        "Random walk (no change)": {"kind": "random_walk", "cols": [], "group": "Benchmark"}})
    bt = run_backtest(hist, methods)
    acc = accuracy_tables(bt)
    loo = pd.concat([leave_one_out(bt, h) for h in HORIZONS], ignore_index=True)
    ojoh = pd.concat([origin_johansen(hist, MODEL_1), origin_johansen(hist, MODEL_2)], ignore_index=True)

    # ---- Report tables 5, A4 and D1 (two years ahead, all forecasts) --------------
    two = acc[(acc.Sample == "All forecasts") & (acc.Horizon == HORIZON)].set_index("Method")
    table5 = two.loc[PRINCIPAL, ["RMSE", "MAE", "MAPE", "Mean error"]].reset_index().rename(columns={"Method": "Model"})
    shortlist = joh[(joh["Price alpha"] < 0) & (joh["Alpha p-value"] < SIG)]["Specification"]
    lookup = {spec_name(MODEL_1): "Model 1", spec_name(MODEL_2): "Model 2"}
    a4 = pd.DataFrame([{"Specification": s, **two.loc[lookup.get(s, s), ["RMSE", "MAE", "MAPE", "Mean error"]].to_dict()}
                       for s in shortlist])
    d1 = []
    for T in [o for o in origins(hist) if o + HORIZON <= SAMPLE_END]:
        sub = bt[(bt.Origin == T) & (bt.Horizon == HORIZON)].set_index("Method")
        row = {"Forecast year": T + HORIZON, "Estimation period": f"{SAMPLE_START}-{T}",
               "Actual (US$/t)": sub.Actual.iloc[0]}
        for m in PRINCIPAL:
            row[m], row[m + " error"] = sub.at[m, "Forecast"], sub.at[m, "Error"]
        row["Model 2 identification"] = sub.at["Model 2", "Identification"]
        d1.append(row)
    d1 = pd.DataFrame(d1)

    # ---- Table 6: final forecasts, each two years ahead ---------------------------
    print("Final forecasts ...")
    t6, final_markup = {m: {} for m in PRINCIPAL}, {}
    for target in FINAL_TARGETS:
        origin = target - HORIZON                          # 2026 <- data to 2024; 2027 <- data to 2025
        train = df.loc[SAMPLE_START:origin]
        mc_target = float(df.at[target, "marginal_cost"])  # projected marginal cost from the cost curve
        final_markup[target] = markup(train)
        t6["Model 1"][target] = vecm_forecast(train, MODEL_1)[HORIZON - 1]
        t6["Model 2"][target] = vecm_forecast(train, MODEL_2)[HORIZON - 1]
        t6["Calibrated Marginal Cost Model"][target] = mc_target * final_markup[target]
        t6["Marginal Cost Model"][target] = mc_target
    table6 = pd.DataFrame([{"Forecasting approach": m, **{str(y): v[y] for y in FINAL_TARGETS}}
                           for m, v in t6.items()])

    # ---- Longer paths to 2035 (beyond the tested two-year horizon) ----------------
    m1_path = vecm_forecast(hist, MODEL_1, LAST_YEAR - SAMPLE_END)
    m2_path = vecm_forecast(hist, MODEL_2, LAST_YEAR - SAMPLE_END)
    paths = []
    for i, y in enumerate(range(SAMPLE_END + 1, LAST_YEAR + 1)):
        m = final_markup.get(y, final_markup[2027])        # 2026 uses data to 2024; later years data to 2025
        paths.append({"Year": y, "Marginal Cost Model": df.at[y, "marginal_cost"],
                      "Calibrated Marginal Cost Model": df.at[y, "marginal_cost"] * m,
                      "Model 1 (path from 2025)": m1_path[i], "Model 2 (path from 2025)": m2_path[i],
                      "Marginal producer": df.at[y, "marginal_asset"]})
    paths = pd.DataFrame(paths)

    # ---- Table E1: demand sensitivity ---------------------------------------------
    e1 = []
    for target in FINAL_TARGETS:
        for shift in (-50, -25, 0, 25, 50):
            d = D["demand"]["total"][target] + shift
            mc, asset = marginal_cost(D, target, demand=d)
            e1.append({"Year": target, "Demand change (Mt)": shift, "Seaborne demand (Mt)": d,
                       "Marginal producer": asset, "Marginal Cost Model (US$/t)": mc,
                       "Calibrated Marginal Cost Model (US$/t)": mc * final_markup[target]})
    e1 = pd.DataFrame(e1)

    # ---- Government forecasts, scored the same way (indicative only) ---------------
    gov = pd.DataFrame([{"Year": y, "Government forecast": f, "Actual": float(hist.at[y, "price"]),
                         "Error": f - float(hist.at[y, "price"]),
                         "Actual above forecast (%)": (float(hist.at[y, "price"]) / f - 1) * 100}
                        for y, f in sorted(D["government"].items()) if y in hist.index])
    g = gov.rename(columns={"Government forecast": "Forecast"})
    gov_acc = pd.DataFrame([{"Method": "Australian Government (indicative, different price basis)", **metrics(g)}])

    # ---- Long-run relationships for Models 1 and 2 (report section 3.1.3) --------
    longrun = diags[diags.Specification.isin([spec_name(MODEL_1), spec_name(MODEL_2)])].dropna(axis=1, how="all")

    # Calibration factor a two-year-ahead forecaster would use for each target
    # year: data to (year - 2), capped at 2025. Used by the dashboard sliders.
    markup_for = {y: markup(df.loc[SAMPLE_START:min(y - HORIZON, SAMPLE_END)]) if y - HORIZON >= SAMPLE_START + 1
                  else markup(hist) for y in range(SAMPLE_START, LAST_YEAR + 1)}

    hash_after = sha256(RAW_DATA)
    checks.loc[len(checks)] = {"Check": "Raw data file unchanged after run",
                               "Result": "Yes" if hash_after == hash_before else "NO - FILE CHANGED"}

    return {"df": df, "hist": hist, "balance": balance, "curves": curves, "checks": checks, "stat": stat,
            "joh": joh, "steps": steps, "diags": diags, "lags": lags, "longrun": longrun,
            "bt": bt, "acc": acc, "loo": loo, "ojoh": ojoh, "table5": table5, "a4": a4, "d1": d1,
            "table6": table6, "paths": paths, "e1": e1, "gov": gov, "gov_acc": gov_acc,
            "final_markup": final_markup, "markup_for": markup_for, "government": D["government"]}


# =============================================================================
# 10. CHARTS (used in the Word document)
# =============================================================================
def chart_cost_curve(R, year):
    """One annual cost curve, drawn in the report's style."""
    rows, info = R["curves"][year], R["df"].loc[year]
    fig, ax = plt.subplots(figsize=(11.69, 8.27))
    x0 = 0.0
    for r in rows:
        ax.bar(x0, r["cost"], width=r["supply_mt"], align="edge", color=COLOURS[r["tier"]],
               edgecolor="white", linewidth=0.3)
        x0 += r["supply_mt"]
    ax.axvline(info.demand, color=COLOURS["demand"], linestyle=":", linewidth=1.5)
    ax.axhline(info.marginal_cost, color=COLOURS["mc"], linewidth=1.5)
    handles = [Patch(color=COLOURS[1], label="Tier 1"), Patch(color=COLOURS[2], label="Tier 2"),
               Patch(color=COLOURS[3], label="Tier 3"),
               Line2D([], [], color=COLOURS["mc"], linewidth=1.5, label="Marginal cost"),
               Line2D([], [], color=COLOURS["demand"], linestyle=":", linewidth=1.5, label="Seaborne demand")]
    top = max([r["cost"] for r in rows] + [info.marginal_cost])
    if year <= SAMPLE_END:
        ax.axhline(info.price, color=COLOURS["price"], linestyle="--", linewidth=1.5)
        handles.insert(3, Line2D([], [], color=COLOURS["price"], linestyle="--", linewidth=1.5,
                                 label="Realised price"))
        top = max(top, info.price)
    step = next(s for s in (10, 20, 25, 50) if top * 1.1 / s <= 9)
    ymax = math.ceil(top * 1.1 / step) * step
    ax.set_xlim(0, math.ceil(max(rows[-1]["cumulative_mt"], info.demand) * 1.07 / 50) * 50)
    ax.set_ylim(0, ymax)
    ax.set_yticks(range(0, int(ymax) + 1, step))
    ax.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    ax.set_xlabel("Cumulative seaborne supply (Mt wet)")
    ax.set_ylabel("Delivered CFR Qingdao cost (US$/t)")
    ax.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, 1.01), ncol=len(handles),
              frameon=False, fontsize=9)
    fig.tight_layout()
    return fig


def chart_backtest(R):
    """Realised prices against each principal model's two-year-ahead forecasts."""
    fig, ax = plt.subplots(figsize=(10, 5.2))
    d1 = R["d1"]
    ax.plot(d1["Forecast year"], d1["Actual (US$/t)"], color="black", linewidth=2.2, label="Realised price")
    styles = dict(zip(PRINCIPAL, ["#2A4879", "#406E6C", "#AB7A5D", "#83525D"]))
    for m, colour in styles.items():
        ax.plot(d1["Forecast year"], d1[m], marker="o", color=colour, label=m)
    gov = R["government"]
    if gov:
        ax.plot(list(gov), list(gov.values()), linestyle="--", color="#8C8C8C",
                label="Australian Government (indicative)")
    ax.set_xticks(list(d1["Forecast year"]))
    ax.set_ylabel("US$/t")
    ax.legend(frameon=False, fontsize=8, ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.08))
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    return fig


def fig_to_stream(fig, dpi=150):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi)
    plt.close(fig)
    buf.seek(0)
    return buf


# =============================================================================
# 11. OUTPUT 1 - WORD RESULTS DOCUMENT
# =============================================================================
def write_docx(R):
    print(f"Writing {OUT_DOCX.name} (this draws 27 charts) ...")
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))

    def table(frame, caption, note=None, fmt=None, size=8):
        """Add a pandas table with a caption. `fmt` maps column -> format string."""
        fmt = fmt or {}
        t = doc.add_table(rows=1, cols=len(frame.columns))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(frame.columns):
            t.rows[0].cells[i].text = str(c)
        for _, row in frame.iterrows():
            cells = t.add_row().cells
            for i, c in enumerate(frame.columns):
                v = row[c]
                if v is None or (isinstance(v, float) and np.isnan(v)):
                    txt = "n.a."
                elif isinstance(v, (bool, np.bool_)):
                    txt = "Yes" if v else "No"
                elif isinstance(v, (int, float, np.floating, np.integer)):
                    txt = fmt.get(c, "{:,.2f}").format(v)
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    txt = str(v)
                cells[i].paragraphs[0].add_run(txt)
        for i, c in enumerate(frame.columns):
            t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size)
        cap = doc.add_paragraph()
        cap.add_run(caption).italic = True
        if note:
            n = doc.add_paragraph(note)
            n.runs[0].font.size = Pt(8)

    p3 = lambda cols: {c: "{:.3f}" for c in cols}

    # --- Title page and summary ---------------------------------------------------
    doc.add_heading("Iron Ore Price Forecasting: Model Results", 0)
    doc.add_paragraph(f"Generated by iron_ore_forecasting.py from {RAW_DATA.name}. All headline results are "
                      "two-year-ahead forecasts on a 62% Fe CFR China basis. Part 1 reproduces the tables in the "
                      "report; Part 2 contains every supporting diagnostic.")
    t5 = R["table5"].set_index("Model")
    t6 = R["table6"].set_index("Forecasting approach")
    doc.add_heading("Summary", 1)
    for line in [
        f"Model 1 (P + MC) has the lowest backtest RMSE (US${t5.at['Model 1', 'RMSE']:.2f}/t) and MAE "
        f"(US${t5.at['Model 1', 'MAE']:.2f}/t), but does not pass cointegration testing.",
        f"Model 2 (P + MC + D) is the only specification that passes both unadjusted Johansen tests and has a "
        f"negative, significant price adjustment. It reduces RMSE by "
        f"{100 * (1 - t5.at['Model 2', 'RMSE'] / t5.at['Marginal Cost Model', 'RMSE']):.0f}% against the "
        f"Marginal Cost Model. No specification passes after the Reinsel-Ahn correction.",
        f"The Calibrated Marginal Cost Model has the lowest MAPE ({t5.at['Calibrated Marginal Cost Model', 'MAPE']:.2f}%).",
        f"2026 forecasts: Model 1 US${t6.at['Model 1', '2026']:.2f}, Model 2 US${t6.at['Model 2', '2026']:.2f}, "
        f"Calibrated US${t6.at['Calibrated Marginal Cost Model', '2026']:.2f}, Marginal Cost Model "
        f"US${t6.at['Marginal Cost Model', '2026']:.2f}/t.",
        f"2027 forecasts: Model 1 US${t6.at['Model 1', '2027']:.2f}, Model 2 US${t6.at['Model 2', '2027']:.2f}, "
        f"Calibrated US${t6.at['Calibrated Marginal Cost Model', '2027']:.2f}, Marginal Cost Model "
        f"US${t6.at['Marginal Cost Model', '2027']:.2f}/t.",
        f"Data checks: {R['checks'].iloc[-1]['Result']} (raw data file unchanged)."]:
        doc.add_paragraph(line, style="List Bullet")

    # --- PART 1: report tables --------------------------------------------------------
    doc.add_heading("Part 1. Results reported in the paper", 1)
    table(R["table5"], "Table 5. Two-year-ahead forecasting accuracy, 2021-2025 (RMSE, MAE, mean error in US$/t; MAPE in %).")
    doc.add_picture(fig_to_stream(chart_backtest(R)), width=Cm(22))
    doc.add_paragraph().add_run("Figure 3. Realised prices and two-year-ahead forecasts, 2021-2025 (US$/t).").italic = True
    table(R["table6"], "Table 6. Forecast iron ore prices (US$/t), 2026-2027.",
          f"All forecasts are two years ahead: 2026 uses data to 2024 (calibration factor "
          f"{R['final_markup'][2026]:.3f}) and 2027 uses data to 2025 ({R['final_markup'][2027]:.3f}). "
          f"2026 year-to-date (to October) average: US${YTD_2026_PRICE}/t, shown for context only.")
    table(R["longrun"][["Specification", "Intercept (long run)", "Long-run MC", "Long-run MC s.e."]
                       + (["Long-run D", "Long-run D s.e."] if "Long-run D" in R["longrun"] else [])
                       + ["Alpha P", "Alpha P p"]],
          "Long-run relationships and price adjustment, Models 1 and 2 (ln P = intercept + coefficients x ln variables).",
          fmt=p3(R["longrun"].columns))

    doc.add_page_break()
    doc.add_heading("Appendix A. Econometric model selection", 1)
    table(R["stat"][["Variable", "Terms", "ADF level p", "ADF diff p", "KPSS level p", "KPSS diff p", "Result",
                     "Proceeds to cointegration testing"]],
          "Table A1. ADF and KPSS stationarity tests, 2010-2025 (p-values).",
          "KPSS p-values are interpolated from tables and bounded at 0.01 and 0.10.",
          fmt=p3(["ADF level p", "ADF diff p", "KPSS level p", "KPSS diff p"]))
    a2 = R["joh"][["Specification", "Trace r=0", "Trace CV r=0", "Corrected trace r=0", "Max-eig r=0",
                   "Max-eig CV r=0", "Corrected max-eig r=0", "Result (unadjusted -> corrected)"]]
    table(a2, "Table A2. Johansen rank-zero tests for all 15 specifications.",
          "Restricted constant, one lagged difference, 14 effective observations. For specifications showing a "
          "relationship, the result also reflects the follow-up tests, which are in Part 2 (Table S1).")
    table(R["joh"][["Specification", "Price alpha", "Alpha p-value", "Alpha assessment"]],
          "Table A3. Price-adjustment coefficients for all 15 VECMs.", fmt=p3(["Price alpha", "Alpha p-value"]))
    table(R["a4"], "Table A4. Two-year-ahead errors for shortlisted models (negative, significant alpha), 2021-2025.")

    doc.add_page_break()
    doc.add_heading("Appendix B. Seaborne supply-demand balance", 1)
    bal = R["balance"].T.reset_index().rename(columns={"index": "Mt wet"})
    for start in range(SAMPLE_START, LAST_YEAR + 1, 9):     # split into readable panels of 9 years
        cols = ["Mt wet"] + list(range(start, min(start + 9, LAST_YEAR + 1)))
        panel = bal[cols].copy()
        panel.columns = [str(c) + ("*" if isinstance(c, int) and c > SAMPLE_END else "") for c in cols]
        table(panel, f"Table B. Global seaborne iron ore supply-demand balance, {start}-{min(start + 8, LAST_YEAR)}.",
              fmt={c: "{:,.0f}" for c in panel.columns}, size=7)
    doc.add_paragraph("* projection. Market balance = change in Tier-3 required + change in visible inventories. "
                      "Marginal cost and price rows are in US$/t (shown rounded).").runs[0].font.size = Pt(8)

    doc.add_page_break()
    doc.add_heading("Appendix C. Annual cost curves", 1)
    for i, year in enumerate(sorted(R["curves"])):
        if i:
            doc.add_page_break()
        doc.add_picture(fig_to_stream(chart_cost_curve(R, year)), width=Cm(23))
        doc.add_paragraph().add_run(f"Figure C{i + 1}. Global seaborne iron ore cost curve (US$/t; Mt wet), {year}"
                                    + (" (projection)" if year > SAMPLE_END else "")
                                    + ". Axis scales vary by year.").italic = True

    doc.add_page_break()
    doc.add_heading("Appendix D. Two-year-ahead backtest forecasts", 1)
    table(R["d1"], "Table D1. Recursive two-year-ahead forecasts and errors (US$/t), 2021-2025.",
          "Errors are forecast minus actual. Model 2's 2021 forecast uses only eight effective observations and is "
          "numerically fragile (statsmodels 0.15.0 gives 145.75).")
    doc.add_heading("Appendix E. Demand sensitivity", 1)
    table(R["e1"], "Table E1. Sensitivity of Marginal Cost Model and Calibrated Marginal Cost Model prices to seaborne demand, 2026-2027.",
          "Each scenario shifts demand while holding supply and costs constant.",
          fmt={"Demand change (Mt)": "{:+,.0f}", "Seaborne demand (Mt)": "{:,.0f}"})

    # --- PART 2: supporting diagnostics ------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 2. Supporting diagnostics (beyond the report)", 1)
    doc.add_paragraph("These tables support the choices made in the report. The dashboard shows the same results interactively.")
    table(R["checks"], "Table S0. Data checks.", size=7)
    table(R["steps"], "Table S1. Johansen tests at every rank step: restricted constant (used), Reinsel-Ahn corrected, "
                      "and conventional unrestricted constant (comparison).", fmt=p3(["Eigenvalue"]), size=6)
    d = R["diags"]
    core = ["Specification", "Alpha P", "Alpha P p", "Largest non-unit root", "Stable system",
            "Ljung-Box (lag 2) p, P eq", "Jarque-Bera p, P eq", "System normality p", "System whiteness (3 lags) p"]
    table(d[core], "Table S2. VECM diagnostics for all 15 specifications (price equation and system tests).",
          "Stable system = largest non-unit root below 1. Ljung-Box tests for remaining autocorrelation; "
          "Jarque-Bera and system normality test whether residuals are normal (p above 0.05 = no problem detected).",
          fmt=p3(core[1:]), size=7)
    granger = ["Specification"] + [c for c in d.columns if c.startswith("Granger")]
    table(d[granger], "Table S3. Granger causality within each VECM: does each variable help predict price? (p-values)",
          fmt=p3(granger[1:]), size=7)
    table(R["lags"], "Table S4. Lag selection by information criterion (lagged differences searched 0-2).",
          "A criterion choosing 2 is at the edge of the search range. The report uses 1 for every model.", fmt={})
    for h in HORIZONS:
        sub = R["acc"][(R["acc"].Horizon == h)].drop(columns="Horizon")
        table(sub, f"Table S5.{h}. {h}-year-ahead accuracy for every model and benchmark, all forecasts and the "
                   f"identified sample.", fmt={"n": "{:.0f}", "Non-identified forecasts": "{:.0f}"}, size=7)
    table(R["loo"], "Table S6. Leave one target year out: RMSE of the principal models, one and two years ahead.",
          fmt={"Horizon": "{:.0f}", "Year left out": "{:.0f}"})
    table(R["ojoh"], "Table S7. Johansen evidence at each forecast origin (Models 1 and 2).",
          fmt={"Data to": "{:.0f}", "Effective obs": "{:.0f}", "Trace rank": "{:.0f}", "Max-eig rank": "{:.0f}",
               "Corrected trace rank": "{:.0f}"})
    table(R["paths"], "Table S8. Paths to 2035. Only 2026-2027 are two-year-ahead forecasts; later years are "
                      "structural scenarios (cost curve) and dynamic VECM paths from 2025, not validated forecasts.")
    table(pd.concat([R["gov"], R["gov_acc"]], ignore_index=True),
          "Table S9. Australian Government two-year-ahead forecasts against realised prices (indicative only: "
          "different price basis).", fmt={"Year": "{:.0f}", "n": "{:.0f}"})
    bt = R["bt"][["Method", "Horizon", "Origin", "Target", "Actual", "Forecast", "Error", "Abs % error", "Identification"]]
    table(bt, "Table S10. Every backtest forecast, every model, one and two years ahead (US$/t).",
          fmt={"Horizon": "{:.0f}", "Origin": "{:.0f}", "Target": "{:.0f}"}, size=6)
    annual = R["df"][["price", "marginal_cost", "marginal_asset", "demand", "total_supply", "assets_producing",
                      "dxy", "bdi", "gdp"]].reset_index()
    annual.columns = ["Year", "Price (US$/t)", "Marginal cost (US$/t)", "Marginal producer", "Seaborne demand (Mt)",
                      "Total supply on curve (Mt)", "Assets producing", "DXY", "BDI", "China real GDP"]
    table(annual, "Table S11. Annual series used by the model (2026-2035 are workbook projections).",
          fmt={"Year": "{:.0f}", "Assets producing": "{:.0f}", "Seaborne demand (Mt)": "{:,.0f}",
               "Total supply on curve (Mt)": "{:,.0f}", "BDI": "{:,.0f}", "China real GDP": "{:,.0f}"}, size=7)
    doc.save(OUT_DOCX)


# =============================================================================
# 12. OUTPUT 2 - INTERACTIVE HTML DASHBOARD (self-contained, works offline)
# =============================================================================
def plotly_library():
    """The charting library is embedded in the HTML so it opens without
    internet. It comes from the 'plotly' Python package."""
    try:
        from plotly.offline import get_plotlyjs
        return "<script>" + get_plotlyjs() + "</script>"
    except ImportError:
        print("  Note: 'plotly' is not installed, so the dashboard will load its charts from the internet.\n"
              "        Run 'pip install plotly' and re-run to make it work offline.")
        return '<script src="https://cdn.plot.ly/plotly-2.35.2.min.js"></script>'


DASHBOARD = r"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Iron Ore Forecasting Dashboard</title>
__PLOTLY__
<style>
:root{--bg:#10131a;--panel:#171c25;--panel2:#1e2531;--ink:#e6e4df;--muted:#97a0ac;--line:#2a313d;
--t1:#4aa3d0;--t2:#6fbf7a;--t3:#e0894f;--px:#f2efe8;--warn:#d8b24a;--good:#6fbf7a}
*{box-sizing:border-box}html,body{margin:0;height:100%}
body{background:var(--bg);color:var(--ink);font:14px -apple-system,BlinkMacSystemFont,"Segoe UI",Helvetica,Arial,sans-serif}
header{padding:10px 16px;border-bottom:1px solid var(--line);background:var(--panel);display:flex;align-items:center;gap:14px;flex-wrap:wrap}
h1{font-size:15px;margin:0;font-weight:600}.sub{font-size:12px;color:var(--muted)}
nav{margin-left:auto;display:flex;gap:6px;flex-wrap:wrap}
nav button{font:inherit;font-size:13px;padding:6px 13px;border:1px solid var(--line);background:var(--panel2);border-radius:6px;cursor:pointer;color:var(--muted)}
nav button:hover{color:var(--ink)}nav button.on{background:var(--t1);border-color:var(--t1);color:#0d1116;font-weight:600}
#main{display:flex;height:calc(100vh - 53px)}
aside{width:280px;flex:none;border-right:1px solid var(--line);background:var(--panel);overflow:auto;padding:12px 14px}
aside h3{font-size:10.5px;text-transform:uppercase;letter-spacing:.07em;color:var(--muted);margin:16px 0 6px;font-weight:600}
aside h3:first-child{margin-top:0}
input[type=search]{width:100%;padding:6px 8px;border:1px solid var(--line);border-radius:5px;font:inherit;background:var(--panel2);color:var(--ink)}
input[type=range]{width:100%;accent-color:var(--t1)}
.lbl{display:flex;justify-content:space-between;font-size:12px;color:var(--muted);margin-top:6px}.lbl b{color:var(--ink)}
.chk{display:flex;align-items:center;gap:7px;font-size:13px;padding:2px 0;cursor:pointer}.chk input{accent-color:var(--t1)}
.dot{width:9px;height:9px;border-radius:50%;flex:none}
.readout{font-size:12px;color:var(--muted);margin-top:5px;line-height:1.55}.readout b{color:var(--ink);font-size:14px}
.scen{background:var(--panel2);border:1px solid var(--line);border-radius:8px;padding:9px 11px;margin-top:8px;font-size:12.5px;line-height:1.6}
.scen .big{font-size:18px;font-weight:600;color:var(--ink)}.up{color:#e2674f}.down{color:var(--good)}
.pill{display:inline-block;padding:3px 10px;border:1px solid var(--line);border-radius:999px;font-size:12px;cursor:pointer;margin:6px 4px 2px 0;background:var(--panel2);color:var(--muted)}
.pill.on{background:var(--t1);color:#0d1116;border-color:var(--t1);font-weight:600}
#panes{flex:1;position:relative;overflow:hidden}.pane{display:none;position:absolute;inset:0;overflow:auto}.pane.on{display:block}
.plotwrap{position:absolute;inset:0 0 40px 0}.foot{position:absolute;left:12px;right:12px;bottom:6px;font-size:11.5px;color:var(--muted)}
.pad{padding:16px 20px 40px}h2{font-size:14px;margin:0 0 8px;font-weight:600}h2.sp{margin-top:26px}
.tw{overflow:auto;max-width:100%}
table{border-collapse:collapse;font-size:12px}
th,td{padding:4px 9px;border-bottom:1px solid var(--line);text-align:right;white-space:nowrap}
th{color:var(--muted);font-weight:600}th:first-child,td:first-child{text-align:left}
td.l,th.l{text-align:left}
table.grid th:first-child,table.grid td:first-child{position:sticky;left:0;background:var(--bg)}
tr.sec td{color:var(--t1);font-weight:600;padding-top:12px}tr.em td{font-weight:600}
.fc{color:#e2674f}.best{color:var(--good);font-weight:600}.bad{color:var(--warn)}
p.note{max-width:1000px;font-size:12.5px;color:var(--muted);line-height:1.6}
.box{background:var(--panel2);border:1px solid var(--line);border-radius:8px;padding:11px 15px;margin:6px 0 12px;font-size:13px;line-height:1.55;max-width:1000px}
.chart{height:430px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:10px;margin-bottom:14px;max-width:1100px}
.card{background:var(--panel2);border:1px solid var(--line);border-radius:8px;padding:10px 12px}.card .k{font-size:11.5px;color:var(--muted)}.card .v{font-size:19px;font-weight:600;margin-top:3px}
#detail{position:absolute;right:16px;top:16px;width:320px;background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:13px 15px;display:none;box-shadow:0 8px 28px rgba(0,0,0,.5);font-size:12.5px;z-index:5}
#detail h4{margin:0 0 9px;font-size:13px;padding-right:16px}.kv{display:grid;grid-template-columns:auto 1fr;gap:3px 12px}.kv .k{color:var(--muted)}.kv .v{text-align:right}
#dclose{position:absolute;right:9px;top:7px;border:0;background:none;font-size:17px;cursor:pointer;color:var(--muted)}
#err{display:none;background:#4a1d1d;border:1px solid #a33;color:#fdd;padding:8px 12px;font-size:13px}
</style></head><body>
<header><h1>Iron ore price forecasting: cost curve + VECM</h1><span class="sub" id="hdr"></span>
<nav id="tabs"><button data-t="curve" class="on">Cost curve</button><button data-t="balance">Balance table</button>
<button data-t="backtest">Backtest</button><button data-t="forecast">Forecasts</button>
<button data-t="selection">Model selection</button><button data-t="diag">Diagnostics</button></nav></header>
<div id="err"></div>
<div id="main">
<aside id="rail">
 <h3>Year</h3><input id="ysl" type="range" min="2010" max="2035" step="1" value="2027"><div class="readout" id="yRead"></div>
 <h3>Demand</h3><input id="dsl" type="range" min="800" max="2200" step="5"><div class="lbl"><span>Seaborne demand</span><b id="dVal"></b></div>
 <h3>Supply change by tier</h3>
 <div class="lbl"><span>Tier 1 (majors)</span><b id="s1Val"></b></div><input id="s1" type="range" min="-30" max="30" step="1" value="0">
 <div class="lbl"><span>Tier 2 (established)</span><b id="s2Val"></b></div><input id="s2" type="range" min="-30" max="30" step="1" value="0">
 <div class="lbl"><span>Tier 3 (swing / junior)</span><b id="s3Val"></b></div><input id="s3" type="range" min="-30" max="30" step="1" value="0">
 <span class="pill" id="reset">reset demand and supply to the workbook</span>
 <div class="scen" id="scen"></div>
 <h3>Search</h3><input id="q" type="search" placeholder="Mine, operator, country...">
 <h3>Tier</h3><div id="tierBox"></div>
 <h3>Country</h3><div id="ctryBox" style="max-height:210px;overflow:auto"></div>
</aside>
<div id="panes">
 <div class="pane on" id="p-curve"><div class="plotwrap" id="curvePlot"></div>
  <div class="foot">Sliders re-read the marginal producer from the cost curve; the calibrated price applies the two-year-ahead calibration factor for that year. They do not re-estimate the VECMs. Click any block for the mine's details.</div>
  <div id="detail"><button id="dclose">&times;</button><div id="dBody"></div></div></div>
 <div class="pane" id="p-balance"><div class="pad" id="balBox"></div></div>
 <div class="pane" id="p-backtest"><div class="pad" id="btBox"></div></div>
 <div class="pane" id="p-forecast"><div class="pad" id="fcBox"></div></div>
 <div class="pane" id="p-selection"><div class="pad" id="selBox"></div></div>
 <div class="pane" id="p-diag"><div class="pad" id="diagBox"></div></div>
</div></div>
<script>
/* ===== data written by the Python script ===== */
const R=__DATA__;

/* ===== helpers ===== */
function report(msg){const e=document.getElementById("err");e.style.display="block";e.textContent+=" "+msg;}
window.onerror=(m)=>report("Dashboard error: "+m);
const HAS=typeof Plotly!=="undefined";
if(!HAS)report("The chart library did not load. Re-run the Python script with plotly installed (pip install plotly). Tables still work.");
const C={t1:"#4aa3d0",t2:"#6fbf7a",t3:"#e0894f",px:"#f2efe8",ink:"#e6e4df",muted:"#97a0ac",line:"#2a313d",warn:"#d8b24a",
 "Model 1":"#b48ce0","Model 2":"#6fbf7a","Calibrated Marginal Cost Model":"#e2674f","Marginal Cost Model":"#4aa3d0",
 "ARIMA (best AIC)":"#97a0ac","Random walk (no change)":"#6c727f"};
const DARK={paper_bgcolor:"rgba(0,0,0,0)",plot_bgcolor:"rgba(0,0,0,0)",font:{color:C.ink,size:12},
 hoverlabel:{bgcolor:"#1e2531",bordercolor:C.line,font:{color:C.ink,size:12},align:"left"}};
const AX={gridcolor:C.line,zerolinecolor:C.line,linecolor:C.line,tickfont:{color:C.muted}};
const ax=(t,o)=>Object.assign({},AX,{title:{text:t,font:{size:12,color:C.muted}}},o||{});
const f=(v,d=0)=>(v===null||v===undefined||Number.isNaN(v))?"n.a.":(typeof v==="number"?v.toLocaleString(undefined,{minimumFractionDigits:d,maximumFractionDigits:d}):String(v));
function plot(id,tr,layout){if(!HAS)return;try{Plotly.react(id,tr,Object.assign({},DARK,layout),{responsive:true,displaylogo:false});}catch(e){report(id+": "+e.message);}}
// decimals for a column, decided by its name
function dec(k){if(/^(Year|Horizon|Origin|Target|n|Variables|Effective obs|Data to|AIC|BIC|HQIC|FPE|Used in report|Year left out|Non-identified forecasts|Assets producing|Rank)$/.test(k)||/rank$/i.test(k)||/lags$/.test(k))return 0;
 if(/( p|p-value|^Eigenvalue|alpha|Alpha|factor|root|Long-run|Intercept|Price eq|Granger|Ljung|Jarque|normality|whiteness|stat)/.test(k))return 3;return 2;}
function tbl(rows,opts={}){if(!rows||!rows.length)return "<p class='note'>No rows.</p>";const ks=opts.cols||Object.keys(rows[0]);
 return `<div class="tw"><table><tr>${ks.map(k=>`<th>${k}</th>`).join("")}</tr>`+rows.map(r=>`<tr${opts.rowClass?` class="${opts.rowClass(r)}"`:""}>`+
  ks.map(k=>`<td class="${typeof r[k]==="string"&&!/^(Yes|No)$/.test(r[k])?"l ":""}${opts.cellClass?opts.cellClass(r,k):""}">${f(r[k],dec(k))}</td>`).join("")+"</tr>").join("")+"</table></div>";}

/* ===== shared lookups ===== */
const CURVE=R.curves, YEARS=Object.keys(CURVE).map(Number).sort((a,b)=>a-b), Y=R.yearInfo;
const S={year:2027,dem:null,adj:{1:0,2:0,3:0},tiers:new Set([1,2,3]),off:new Set(),q:""};
document.getElementById("hdr").textContent=`Source: ${R.source} (unchanged: ${R.unchanged}) · estimation 2010-2025 · two-year-ahead forecasts, 62% Fe CFR China basis`;

/* ===== cost curve controls ===== */
const ctry=[...new Set(YEARS.flatMap(y=>CURVE[y].map(r=>r.country||"(not stated)")))].sort();
[[1,"Tier 1: majors",C.t1],[2,"Tier 2: established",C.t2],[3,"Tier 3: swing / junior",C.t3]].forEach(([t,l,c])=>{
 const e=document.createElement("label");e.className="chk";const i=document.createElement("input");i.type="checkbox";i.checked=true;
 i.onchange=()=>{i.checked?S.tiers.add(t):S.tiers.delete(t);drawCurve()};const d=document.createElement("span");d.className="dot";d.style.background=c;
 e.append(i,d,document.createTextNode(l));document.getElementById("tierBox").append(e)});
ctry.forEach(c=>{const e=document.createElement("label");e.className="chk";const i=document.createElement("input");i.type="checkbox";i.checked=true;
 i.onchange=()=>{i.checked?S.off.delete(c):S.off.add(c);drawCurve()};e.append(i,document.createTextNode(c));document.getElementById("ctryBox").append(e)});
const ysl=document.getElementById("ysl"),dsl=document.getElementById("dsl");
const baseDem=y=>Y[y].demand;
ysl.oninput=()=>{S.year=+ysl.value;S.dem=null;dsl.value=Math.round(baseDem(S.year));drawCurve()};
dsl.oninput=()=>{S.dem=+dsl.value;drawCurve()};
[1,2,3].forEach(t=>{document.getElementById("s"+t).oninput=e=>{S.adj[t]=+e.target.value;drawCurve()}});
document.getElementById("reset").onclick=()=>{S.dem=null;S.adj={1:0,2:0,3:0};[1,2,3].forEach(t=>document.getElementById("s"+t).value=0);dsl.value=Math.round(baseDem(S.year));drawCurve()};
document.getElementById("q").oninput=e=>{S.q=e.target.value.toLowerCase();drawCurve()};
document.getElementById("dclose").onclick=()=>document.getElementById("detail").style.display="none";

/* ===== 1. cost curve (the interactive part) ===== */
let shown=[];
function marginalOf(rows,D){let cum=0;for(const r of rows){cum+=r.supply_mt;if(cum>=D)return r;}return null;}
function drawCurve(){try{
 const D=S.dem!=null?S.dem:baseDem(S.year);
 // apply the tier supply sliders, then the filters, then re-stack the curve
 const rows=CURVE[S.year].map(r=>Object.assign({},r,{supply_mt:r.supply_mt*(1+S.adj[r.tier]/100)}))
  .filter(r=>S.tiers.has(r.tier)&&!S.off.has(r.country||"(not stated)"));
 let cum=0,mi=-1;shown=rows.map(r=>{const p=cum;cum+=r.supply_mt;return Object.assign({},r,{x0:p,cum:cum})});
 for(let i=0;i<shown.length;i++){if(shown[i].cum>=D){mi=i;break}}
 const hit=r=>S.q&&(r.asset+" "+r.operator+" "+r.country).toLowerCase().includes(S.q);
 const tr=[1,2,3].map(t=>{const g=shown.map((r,i)=>[r,i]).filter(([r])=>r.tier===t);
  return {type:"bar",name:["Tier 1","Tier 2","Tier 3"][t-1],x:g.map(([r])=>r.x0+r.supply_mt/2),y:g.map(([r])=>r.cost),
   width:g.map(([r])=>r.supply_mt),customdata:g.map(([,i])=>i),textposition:"none",opacity:.92,
   marker:{color:g.map(([r,i])=>i===mi?"#ffffff":[C.t1,C.t2,C.t3][t-1]),line:{width:g.map(([r])=>hit(r)?2:0),color:"#ffffff"}},
   hovertext:g.map(([r])=>`<b>${r.asset}</b><br>${r.operator||"n.a."} · ${r.country||"n.a."}<br>Supply ${f(r.supply_mt,1)} Mt (cumulative ${f(r.cum,0)} Mt)<br>Delivered cost <b>US$${f(r.cost,2)}/t</b> · Tier ${r.tier}`),
   hovertemplate:"%{hovertext}<extra></extra>"}});
 const m=mi>=0?shown[mi]:null, info=Y[S.year];
 const sh=[{type:"line",x0:D,x1:D,y0:0,y1:1,yref:"paper",line:{color:C.ink,width:1.4,dash:"dot"}}];
 const an=[{x:D,y:1,yref:"paper",yanchor:"bottom",showarrow:false,text:`Demand ${f(D,0)} Mt${S.dem!=null?" (scenario)":""}`,font:{size:11}}];
 if(m){sh.push({type:"line",x0:0,x1:1,xref:"paper",y0:m.cost,y1:m.cost,line:{color:C.t1,width:1.4}});
  an.push({x:0,xref:"paper",y:m.cost,xanchor:"left",yanchor:"bottom",showarrow:false,text:`Marginal cost ${f(m.cost,2)}`,font:{size:11,color:C.t1}})}
 else an.push({x:.5,xref:"paper",y:.5,yref:"paper",showarrow:false,text:"Supply shown does not reach demand",font:{size:13,color:C.warn}});
 if(info.price!=null){sh.push({type:"line",x0:0,x1:1,xref:"paper",y0:info.price,y1:info.price,line:{color:C.px,width:1.5,dash:"dash"}});
  an.push({x:1,xref:"paper",y:info.price,xanchor:"right",yanchor:"bottom",showarrow:false,text:"Realised price "+f(info.price,2),font:{size:11,color:C.px}})}
 plot("curvePlot",tr,{margin:{l:62,r:20,t:40,b:48},bargap:0,barmode:"overlay",
  title:{text:`${S.year} delivered cost curve · ${shown.length} assets${S.year>2025?" · workbook projection":""}`,font:{size:13},x:0,xanchor:"left"},
  xaxis:ax("Cumulative seaborne supply (Mt wet)",{rangemode:"tozero"}),yaxis:ax("Delivered CFR Qingdao cost (US$/t)",{rangemode:"tozero"}),
  shapes:sh,annotations:an,legend:{orientation:"h",x:1,xanchor:"right",y:1.02,yanchor:"bottom",font:{size:11,color:C.muted}}});
 // readouts
 const nb=i=>shown[i]?`${shown[i].asset} (${f(shown[i].cost,1)})`:"n.a.";
 document.getElementById("yRead").innerHTML=`<b>${S.year}</b> ${S.year>2025?"projection":"history"}<br>Workbook demand ${f(info.demand,0)} Mt; marginal producer ${info.asset}`;
 document.getElementById("dVal").textContent=f(D,0)+" Mt"+(S.dem!=null?" *":"");
 [1,2,3].forEach(t=>document.getElementById("s"+t+"Val").textContent=(S.adj[t]>0?"+":"")+S.adj[t]+"%");
 const mk=R.markupFor[S.year], baseMC=info.mc, calc=m?m.cost*mk:null, baseCal=baseMC*mk;
 const chg=(a,b)=>a==null?"":`<span class="${a>b?"up":a<b?"down":""}">(${a>=b?"+":""}${f(a-b,2)} vs workbook)</span>`;
 document.getElementById("scen").innerHTML=m?
  `Marginal producer: <b>${m.asset}</b><br>Below: ${nb(mi-1)}<br>Above: ${nb(mi+1)}<br>`+
  `Marginal cost <span class="big">US$${f(m.cost,2)}</span> ${chg(m.cost,baseMC)}<br>`+
  `Calibrated price <span class="big">US$${f(calc,2)}</span> ${chg(calc,baseCal)}<br><span style="color:var(--muted)">Calibration factor ${f(mk,4)} (${R.markupNote[S.year]})</span>`+
  (info.price!=null?`<br>Realised price US$${f(info.price,2)}`:""):`<span class="bad">Supply shown does not reach demand: no marginal producer.</span>`;
 // click a block for the mine's details
 const el=document.getElementById("curvePlot");if(el.removeAllListeners)el.removeAllListeners("plotly_click");
 if(el.on)el.on("plotly_click",ev=>{const i=ev.points[0].customdata;if(i==null)return;const r=shown[i];
  document.getElementById("dBody").innerHTML=`<h4>${r.asset}</h4><div class="kv">`+[["Operator",r.operator],["Owner",r.owner],["Country",r.country],
  ["Region",r.region],["Port",r.port],["Status",r.status],["Tier",r.tier],["Grade",f(r.grade,1)+"% Fe"],["Moisture",f(r.moisture,1)+"%"],
  ["Royalty",f(r.royalty,1)+"%"],["Supply (as shown)",f(r.supply_mt,1)+" Mt wet"],["FOB cost",f(r.fob,2)+" US$/wmt"],["Freight",f(r.freight,2)+" US$/wmt"],
  ["Delivered CFR cost",f(r.cost,2)+" US$/t"],["Calculation","(FOB/(1-royalty)+freight)/(1-moisture)"]]
  .map(([k,v])=>`<div class="k">${k}</div><div class="v">${v||"n.a."}</div>`).join("")+"</div>";document.getElementById("detail").style.display="block"});
}catch(e){report("cost curve: "+e.message);}}

/* ===== 2. balance table ===== */
function drawBalance(){
 const B=R.balance, ks=Object.keys(B[0]).filter(k=>k!=="Year");
 const sec={"China seaborne demand":"Seaborne demand","Tier-1 supply":"Seaborne supply","Market balance":"Market balance and price"};
 const bold=new Set(["China seaborne demand","Total seaborne demand","Tier-1 supply","Market balance","62% Fe price, realised (US$/t)"]);
 let h=`<h2>Seaborne supply-demand balance (Mt wet; prices US$/t)</h2><div class="tw"><table class="grid"><tr><th>Mt wet</th>`+B.map(b=>`<th class="${b.Year>2025?"fc":""}">${b.Year}${b.Year>2025?"*":""}</th>`).join("")+"</tr>";
 ks.forEach(k=>{if(sec[k])h+=`<tr class="sec"><td colspan="${B.length+1}">${sec[k]}</td></tr>`;
  h+=`<tr${bold.has(k)?' class="em"':""}><td>${k}</td>`+B.map(b=>{const v=b[k];const t=v==null?"n.a.":(Math.abs(v)<0.5&&v!==0?"0":v===0?"-":f(v,k.includes("US$")?2:0));return `<td class="${b.Year>2025?"fc":""}">${t}</td>`}).join("")+"</tr>"});
 h+=`</table></div><p class="note">* projection. Rio Tinto includes IOC; Germano is split 50/50 between BHP and Vale; Tier-2 excludes IOC. Tier-3 required = demand - named producers - Tier-2 (negative = Tier 1-2 supply alone exceeds demand). Market balance = change in Tier-3 required + change in visible inventories. A dash means zero.</p>`;
 document.getElementById("balBox").innerHTML=h;}

/* ===== 3. backtest ===== */
let H=2;
function drawBacktest(){
 const rec=R.bt.filter(r=>r.Horizon===H), tg=[...new Set(rec.map(r=>r.Target))].sort();
 const accAll=R.acc.filter(a=>a.Horizon===H&&a.Sample==="All forecasts"), accId=R.acc.filter(a=>a.Horizon===H&&a.Sample!=="All forecasts");
 const best=(rows,k)=>Math.min(...rows.filter(a=>a.Group==="Principal").map(a=>a[k]));
 const accT=rows=>{const b={RMSE:best(rows,"RMSE"),MAE:best(rows,"MAE"),MAPE:best(rows,"MAPE")};
  return tbl(rows,{cols:["Method","Group","n","RMSE","MAE","MAPE","Mean error","Non-identified forecasts","Reliable"],
   cellClass:(r,k)=>(b[k]!==undefined&&r[k]===b[k])?"best":(k==="Reliable"&&r[k]==="No")?"bad":""});};
 const t5=Object.fromEntries(R.table5.map(r=>[r.Model,r]));
 let h=`<h2>Recursive backtest (pseudo-out-of-sample)</h2>
 <div class="cards"><div class="card"><div class="k">Lowest 2-year RMSE</div><div class="v">Model 1: ${f(t5["Model 1"].RMSE,2)}</div></div>
 <div class="card"><div class="k">Lowest 2-year MAPE</div><div class="v">Calibrated: ${f(t5["Calibrated Marginal Cost Model"].MAPE,2)}%</div></div>
 <div class="card"><div class="k">Model 2 RMSE vs Marginal Cost Model</div><div class="v">${f(100*(1-t5["Model 2"].RMSE/t5["Marginal Cost Model"].RMSE),0)}% lower</div></div></div>
 <div class="box">At each origin every model is re-estimated on data up to that year only, then forecasts one and two years ahead. VECMs forecast their own marginal cost and demand; the cost-curve models use the marginal cost known at the origin. <b>Hollow markers</b> are forecasts from a VECM that is <span class="bad">not identified</span> at that origin (too few observations for a unique long-run relationship). The <i>identified sample</i> drops those target years for every model.</div>
 <span class="pill${H===1?" on":""}" data-h="1">1 year ahead</span><span class="pill${H===2?" on":""}" data-h="2">2 years ahead (the research question)</span>
 <div class="chart" id="btPlot"></div>
 <h2 class="sp">Accuracy, all forecasts (${H}-year)</h2>${accT(accAll)}
 <h2 class="sp">Accuracy, identified sample (${H}-year)</h2>${accT(accId)}
 <h2 class="sp">Leave one target year out (${H}-year RMSE)</h2>${tbl(R.loo.filter(r=>r.Horizon===H))}
 <h2 class="sp">Every forecast, principal models (${H}-year)</h2>${tbl(rec.filter(r=>r.Group==="Principal"),{cols:["Method","Origin","Target","Estimation period","Effective obs","Actual","Forecast","Error","Abs % error","Identification"],cellClass:(r,k)=>k==="Identification"&&r[k]==="Not identified"?"bad":""})}
 <h2 class="sp">Every forecast, all other models (${H}-year)</h2>${tbl(rec.filter(r=>r.Group!=="Principal"),{cols:["Method","Group","Origin","Target","Actual","Forecast","Error","Identification"],cellClass:(r,k)=>k==="Identification"&&r[k]==="Not identified"?"bad":""})}`;
 document.getElementById("btBox").innerHTML=h;
 document.querySelectorAll("#btBox .pill").forEach(p=>p.onclick=()=>{H=+p.dataset.h;drawBacktest()});
 const actual=tg.map(t=>rec.find(r=>r.Target===t).Actual);
 const tr=[{x:tg,y:actual,name:"Realised price",mode:"lines+markers",line:{color:C.px,width:2.6},marker:{size:7}}];
 ["Model 1","Model 2","Calibrated Marginal Cost Model","Marginal Cost Model","ARIMA (best AIC)","Random walk (no change)"].forEach(m=>{
  const rr=tg.map(t=>rec.find(r=>r.Target===t&&r.Method===m));const bench=m.startsWith("ARIMA")||m.startsWith("Random");
  tr.push({x:tg,y:rr.map(r=>r.Forecast),name:m,mode:"lines+markers",visible:bench?"legendonly":true,line:{color:C[m],width:bench?1.3:2,dash:bench?"dot":"solid"},
   marker:{size:8,symbol:rr.map(r=>r.Identification==="Not identified"?"circle-open":"circle"),line:{width:2,color:C[m]}},
   hovertext:rr.map(r=>`${m}<br>data to ${r.Origin} → ${r.Target}<br>forecast ${f(r.Forecast,2)} · actual ${f(r.Actual,2)}<br>error ${f(r.Error,2)}<br>${r.Identification}`),hovertemplate:"%{hovertext}<extra></extra>"})});
 if(H===2&&R.gov.length)tr.push({x:R.gov.map(g=>g.Year),y:R.gov.map(g=>g["Government forecast"]),name:"Australian Government (indicative)",mode:"lines+markers",line:{color:"#8C8C8C",dash:"dash"}});
 plot("btPlot",tr,{margin:{l:58,r:20,t:14,b:46},xaxis:ax("Target year",{dtick:1}),yaxis:ax("US$/t"),legend:{orientation:"h",y:-0.2,font:{size:11,color:C.muted}}});}

/* ===== 4. forecasts ===== */
function drawForecast(){
 const t6=Object.fromEntries(R.table6.map(r=>[r["Forecasting approach"],r]));
 let h=`<h2>Final forecasts (each made two years ahead)</h2>
 <div class="cards">${R.table6.map(r=>`<div class="card"><div class="k">${r["Forecasting approach"]}</div><div class="v">${f(r["2026"],2)} → ${f(r["2027"],2)}</div></div>`).join("")}</div>
 <div class="box">2026 forecasts use data to 2024 (calibration factor ${f(R.finalMarkup["2026"],3)}); 2027 forecasts use data to 2025 (${f(R.finalMarkup["2027"],3)}). Solid points are the two-year-ahead forecasts. Dotted lines after 2027 are <b>scenarios</b>: the cost-curve models follow the workbook projection, and the VECM paths run on from 2025. Neither is a tested forecast. The 2026 year-to-date average (US$${f(R.ytd,2)}/t) is shown for context only.</div>
 <div class="chart" id="fcPlot"></div>
 <h2 class="sp">Table 6. Final forecasts (US$/t)</h2>${tbl(R.table6,{cols:["Forecasting approach","2026","2027"]})}
 <h2 class="sp">Demand sensitivity (Table E1)</h2>${tbl(R.e1)}
 <h2 class="sp">Paths to 2035 (only 2026-2027 are two-year-ahead forecasts)</h2>${tbl(R.paths)}`;
 document.getElementById("fcBox").innerHTML=h;
 const hy=YEARS.filter(y=>y<=2025), py=YEARS.filter(y=>y>2025);
 const P=k=>R.paths.map(p=>p[k]);
 const tr=[{x:hy,y:hy.map(y=>Y[y].price),name:"Realised price",mode:"lines+markers",line:{color:C.px,width:2.6},marker:{size:5}},
  {x:YEARS,y:YEARS.map(y=>Y[y].mc),name:"Marginal cost (history + projection)",mode:"lines",line:{color:C["Marginal Cost Model"],width:1.6}}];
 ["Model 1","Model 2","Calibrated Marginal Cost Model"].forEach(m=>{
  tr.push({x:[2025,2026,2027],y:[Y[2025].price,t6[m]["2026"],t6[m]["2027"]],name:m+" (two-year-ahead)",mode:"lines+markers",line:{color:C[m],width:2.4},marker:{size:7}});});
 tr.push({x:py.filter(y=>y>=2027),y:R.paths.filter(p=>p.Year>=2027).map(p=>p["Calibrated Marginal Cost Model"]),name:"Calibrated: structural scenario",mode:"lines",line:{color:C["Calibrated Marginal Cost Model"],width:1.3,dash:"dot"}});
 tr.push({x:py,y:P("Model 2 (path from 2025)"),name:"Model 2: path from 2025",mode:"lines",visible:"legendonly",line:{color:C["Model 2"],width:1.3,dash:"dot"}});
 tr.push({x:py,y:P("Model 1 (path from 2025)"),name:"Model 1: path from 2025",mode:"lines",visible:"legendonly",line:{color:C["Model 1"],width:1.3,dash:"dot"}});
 tr.push({x:[2026],y:[R.ytd],name:"2026 YTD average (context)",mode:"markers",marker:{symbol:"diamond",size:11,color:"#fff"}});
 plot("fcPlot",tr,{margin:{l:58,r:20,t:14,b:46},xaxis:ax("Year",{dtick:5}),yaxis:ax("US$/t",{rangemode:"tozero"}),legend:{orientation:"h",y:-0.2,font:{size:11,color:C.muted}},
  shapes:[{type:"line",x0:2025.5,x1:2025.5,y0:0,y1:1,yref:"paper",line:{color:C.muted,width:1,dash:"dot"}}]});}

/* ===== 5. model selection ===== */
function drawSelection(){
 document.getElementById("selBox").innerHTML=`<h2>Step 1. Stationarity (Table A1)</h2>
 <div class="box">A variable goes forward if it is I(1): non-stationary in levels and stationary after differencing. China real GDP fails because its first difference is not stationary.</div>
 ${tbl(R.stat,{cellClass:(r,k)=>k==="Proceeds to cointegration testing"?(r[k]==="Yes"?"best":"bad"):""})}
 <h2 class="sp">Step 2. Johansen cointegration and price adjustment, all 15 specifications (Tables A2-A3)</h2>
 <div class="box">"Passes" means the unadjusted trace and maximum-eigenvalue tests both find exactly one relationship. Model 2 is the only specification that passes and has a negative, significant price adjustment (alpha). No specification passes after the Reinsel-Ahn correction.</div>
 ${tbl(R.joh,{cellClass:(r,k)=>(k==="Passes cointegration (unadjusted)"&&r[k]==="Yes")||(k==="Alpha assessment"&&r[k]==="Negative; significant")?"best":""})}
 <h2 class="sp">Step 3. Shortlisted models (negative, significant alpha): two-year accuracy (Table A4)</h2>${tbl(R.a4)}`;}

/* ===== 6. diagnostics ===== */
function drawDiag(){
 document.getElementById("diagBox").innerHTML=`<h2>Data checks</h2>${tbl(R.checks)}
 <h2 class="sp">Johansen tests at every rank step</h2><div class="box">Restricted constant (used in the report), with the Reinsel-Ahn correction, and the conventional unrestricted-constant test for comparison.</div>${tbl(R.steps)}
 <h2 class="sp">VECM diagnostics, all 15 specifications</h2><div class="box">Long-run coefficients, adjustment speeds and short-run terms; stability (largest non-unit root below 1 = stable); residual tests (p above 0.05 = no problem detected); Granger causality (p below 0.05 = the variable helps predict price).</div>${tbl(R.diags)}
 <h2 class="sp">Lag selection</h2>${tbl(R.lags)}
 <h2 class="sp">Cointegration evidence as the sample grows (Models 1 and 2)</h2>${tbl(R.ojoh)}
 <h2 class="sp">Australian Government two-year-ahead forecasts (indicative: different price basis)</h2>${tbl(R.govTable)}`;}

/* ===== tabs ===== */
const DRAW={curve:drawCurve,balance:drawBalance,backtest:drawBacktest,forecast:drawForecast,selection:drawSelection,diag:drawDiag};
document.querySelectorAll("#tabs button").forEach(b=>b.onclick=()=>{
 document.querySelectorAll("#tabs button").forEach(x=>x.classList.toggle("on",x===b));const t=b.dataset.t;
 document.querySelectorAll(".pane").forEach(p=>p.classList.toggle("on",p.id==="p-"+t));
 document.getElementById("rail").style.display=t==="curve"?"block":"none";
 try{DRAW[t]();}catch(e){report(t+": "+e.message);}});
dsl.value=Math.round(baseDem(S.year));drawCurve();
</script></body></html>"""


def to_plain(obj):
    """Convert numpy / pandas values into plain Python so they can be written
    as JavaScript. NaN and infinity become null; True/False become 'Yes'/'No'."""
    if isinstance(obj, dict):
        return {str(k): to_plain(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [to_plain(v) for v in obj]
    if isinstance(obj, pd.DataFrame):
        return [to_plain(r) for r in obj.to_dict("records")]
    if isinstance(obj, (bool, np.bool_)):
        return "Yes" if obj else "No"
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, (float, np.floating)):
        return None if not np.isfinite(obj) else float(obj)
    return obj


def write_html(R):
    print(f"Writing {OUT_HTML.name} ...")
    df = R["df"]
    data = to_plain({
        "source": RAW_DATA.name, "unchanged": R["checks"].iloc[-1]["Result"],
        "curves": {y: rows for y, rows in R["curves"].items()},
        "yearInfo": {y: {"demand": df.at[y, "demand"], "mc": df.at[y, "marginal_cost"],
                         "asset": df.at[y, "marginal_asset"],
                         "price": df.at[y, "price"] if y <= SAMPLE_END else None} for y in df.index},
        "markupFor": R["markup_for"],
        "markupNote": {y: (f"data to {min(y - HORIZON, SAMPLE_END)}" if y - HORIZON >= SAMPLE_START + 1
                           else "full 2010-2025 sample: too few earlier years") for y in df.index}, "finalMarkup": R["final_markup"], "ytd": YTD_2026_PRICE,
        "balance": R["balance"].reset_index(), "bt": R["bt"], "acc": R["acc"], "loo": R["loo"],
        "table5": R["table5"], "table6": R["table6"], "e1": R["e1"], "paths": R["paths"], "a4": R["a4"],
        "stat": R["stat"].drop(columns="col"), "joh": R["joh"].drop(columns="cols"), "steps": R["steps"],
        "diags": R["diags"], "lags": R["lags"], "ojoh": R["ojoh"], "checks": R["checks"],
        "gov": R["gov"], "govTable": pd.concat([R["gov"], R["gov_acc"]], ignore_index=True)})
    html = DASHBOARD.replace("__PLOTLY__", plotly_library())
    html = html.replace("__DATA__", json.dumps(data, allow_nan=False).replace("</", "<\\/"))
    OUT_HTML.write_text(html, encoding="utf-8")


# =============================================================================
# 13. MAIN
# =============================================================================
if __name__ == "__main__":
    if not RAW_DATA.exists():
        sys.exit(f"Cannot find {RAW_DATA}. Put NPSC3001_RawData.xlsx next to this script "
                 "or pass its path: python iron_ore_forecasting.py path/to/file.xlsx")
    results = run()
    write_docx(results)
    write_html(results)

    # Print the headline numbers so you can check them against the report
    print("\nTable 5 - two-year-ahead accuracy, 2021-2025")
    print(results["table5"].round(2).to_string(index=False))
    print("\nTable 6 - final forecasts (US$/t)")
    print(results["table6"].round(2).to_string(index=False))
    print("\nData check:", results["checks"].iloc[-1]["Check"], "->", results["checks"].iloc[-1]["Result"])
    print(f"Done. Files written to {HERE}")
